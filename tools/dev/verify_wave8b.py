"""End-to-end check of Wave 8 on a development machine (demo mode, mock HOSxP only).

Starts the API (``serve --demo``) and the user interface on spare ports, then, through the
user interface as a browser would:

* opens the nine reports (D-35) and downloads each as Excel for a Plan Officer;
* checks the requester scope (own department only, no audit report);
* checks the times in the Excel file are hospital time (UTC+7);
* checks every download wrote one REPORT_EXPORTED audit row and viewing wrote none;
* checks the demo banner (D-26) on every Wave 8 page;
* checks the database schema is at the latest migration.

Refuses to run unless PPR_DATABASE_URL names a development database (``*_dev``) and the
HOSxP mode is mock: demo mode writes demo data into that database.

    .venv\\Scripts\\python tools\\dev\\verify_wave8b.py
"""

from __future__ import annotations

import http.cookiejar
import os
import re
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import UTC, datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "backend" / "src"))

from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import create_engine, text  # noqa: E402

from ppr.config.settings import Settings, database_name, redact_url  # noqa: E402

API_PORT, UI_PORT = 8765, 3765
UI = f"http://127.0.0.1:{UI_PORT}"
PASSWORD = "demo1234"
BANGKOK = timezone(timedelta(hours=7))
REPORTS = [
    "PLAN_SUMMARY",
    "PLAN_REMAINING",
    "DEPARTMENT_UTILIZATION",
    "CATEGORY_UTILIZATION",
    "PPR_REGISTER",
    "PENDING_PPR",
    "CANCELLED_PPR",
    "PR_CHANGED",
    "AUDIT",
]
results: list[tuple[bool, str]] = []


def check(ok: bool, what: str) -> None:
    results.append((ok, what))
    print(("PASS  " if ok else "FAIL  ") + what, flush=True)


class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *args, **kwargs):  # type: ignore[no-untyped-def]
        return None


def browser(user: str):  # type: ignore[no-untyped-def]
    jar = http.cookiejar.CookieJar()
    op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar), NoRedirect)

    def get(path: str) -> tuple[int, bytes, dict[str, str]]:
        try:
            r = op.open(UI + path, timeout=60)
            return r.status, r.read(), dict(r.headers)
        except urllib.error.HTTPError as e:
            return e.code, e.read(), dict(e.headers)

    _, page, _ = get("/login")
    csrf = re.search(rb'name="_csrf" value="([^"]+)"', page)
    assert csrf, "no login form"
    form = urllib.parse.urlencode(
        {"_csrf": csrf.group(1).decode(), "username": user, "password": PASSWORD}
    ).encode()
    try:
        op.open(urllib.request.Request(UI + "/login", data=form), timeout=60)
    except urllib.error.HTTPError as e:
        if e.code != 303:
            raise
    return get


def wait(url: str, seconds: int = 60) -> bool:
    end = time.time() + seconds
    while time.time() < end:
        try:
            urllib.request.urlopen(url, timeout=3)
            return True
        except Exception:
            time.sleep(1)
    return False


def table(body: bytes) -> bytes:
    """Only the report table (the filter lists name every department)."""
    start, end = body.find(b"<tbody>"), body.rfind(b"</tbody>")
    return body[start:end] if start >= 0 and end > start else b""


def departments(body: bytes) -> list[str]:
    """The department column (4th) of each row of the report table."""
    out = []
    for row in re.findall(rb"<tr>(.*?)</tr>", table(body), re.S):
        cells = re.findall(rb"<td[^>]*>(.*?)</td>", row, re.S)
        if len(cells) >= 4:
            out.append(cells[3].decode().strip())
    return out


def exports(engine) -> int:  # type: ignore[no-untyped-def]
    with engine.connect() as c:
        return int(
            c.execute(
                text("SELECT count(*) FROM audit_log WHERE action = 'REPORT_EXPORTED'")
            ).scalar_one()
        )


def main() -> int:
    settings = Settings.load()
    url = settings.require_database_url()
    if not database_name(url).lower().endswith("_dev") or settings.hosxp_mode != "mock":
        print(
            f"REFUSED: needs a *_dev database and PPR_HOSXP_MODE=mock "
            f"(got {redact_url(url)}, mode {settings.hosxp_mode})"
        )
        return 2
    py = sys.executable
    subprocess.run([py, "-m", "ppr.cli", "migrate"], cwd=REPO / "backend", check=True)
    engine = create_engine(url)
    with engine.connect() as c:
        head = c.execute(text("SELECT version_num FROM alembic_version")).scalar_one()
    check(head == "0007", f"database schema at the latest migration ({head})")

    env = {**os.environ, "PPR_API_URL": f"http://127.0.0.1:{API_PORT}", "PPR_UI_PORT": str(UI_PORT)}
    api = subprocess.Popen(
        [py, "-m", "ppr.cli", "serve", "--demo", "--port", str(API_PORT)],
        cwd=REPO / "backend",
        env=env,
    )
    ui = subprocess.Popen(["node", str(REPO / "frontend" / "src" / "server.js")], env=env)
    try:
        if not (wait(f"http://127.0.0.1:{API_PORT}/api/health") and wait(UI + "/login")):
            check(False, "servers started")
            return 1
        officer = browser("demo_plan")
        requester = browser("demo_req_ent")

        before = exports(engine)
        for path in ("/dashboard", "/alerts", "/plans", "/reports"):
            st, body, _ = officer(path)
            check(st == 200 and b'class="demo-banner"' in body, f"{path}: opens with demo banner")
        downloads = 0
        for code in REPORTS:
            st, body, _ = officer(f"/reports/{code}?go=1")
            check(
                st == 200 and b'class="demo-banner"' in body,
                f"report {code}: opens with demo banner",
            )
            st, data, headers = officer(f"/reports/{code}/xlsx")
            ctype = headers.get("Content-Type", "")
            ok = st == 200 and "spreadsheetml" in ctype and data[:2] == b"PK"
            check(ok, f"report {code}: Excel downloads")
            if not ok:
                continue
            downloads += 1
            ws = load_workbook(BytesIO(data)).active
            made = ws.cell(row=2, column=2).value
            now_bkk = datetime.now(UTC).astimezone(BANGKOK).replace(tzinfo=None)
            check(
                isinstance(made, datetime) and abs((now_bkk - made).total_seconds()) < 600,
                f"report {code}: time in the file is Thai time ({made})",
            )
        after_views = exports(engine)
        check(
            after_views - before == downloads,
            f"{downloads} downloads = {after_views - before} REPORT_EXPORTED rows",
        )

        st, _, _ = officer("/reports/PPR_REGISTER?go=1")
        check(exports(engine) == after_views, "viewing a report on screen writes no audit row")

        st, body, _ = requester("/reports")
        check(
            st == 200 and "ประวัติการตรวจสอบ (Audit)".encode() not in body,
            "requester: audit report not offered",
        )
        st, _, _ = requester("/reports/AUDIT?go=1")
        check(st == 403, f"requester: audit report refused ({st})")
        mark = exports(engine)
        st, _, headers = requester("/reports/AUDIT/xlsx")
        check(
            st == 303 and exports(engine) == mark, "requester: audit file refused, nothing audited"
        )
        # Register: the requester's rows must be exactly the officer's rows of Mock ENT.
        st, body, _ = requester("/reports/PPR_REGISTER?go=1")
        mine = departments(body)
        _, all_body, _ = officer("/reports/PPR_REGISTER?go=1")
        everyone = departments(all_body)
        check(
            st == 200 and set(mine) <= {"Mock ENT"} and len(mine) == everyone.count("Mock ENT"),
            f"requester: register shows own department only (requester sees {len(mine)} row(s): "
            f"{sorted(set(mine))}; all departments: {len(everyone)} row(s) "
            f"{sorted(set(everyone))})",
        )
        st, body, _ = requester("/reports/PLAN_REMAINING?go=1")
        rows = table(body)
        check(
            st == 200 and b"Mock ENT" in rows and b"Mock OR" not in rows,
            "requester: plan report shows own department only",
        )
        st, body, _ = officer("/reports/PLAN_REMAINING?go=1")
        check(b"Mock OR" in table(body), "plan officer: plan report shows every department")
    finally:
        for p in (ui, api):
            p.terminate()
            try:
                p.wait(timeout=15)
            except subprocess.TimeoutExpired:
                p.kill()
        engine.dispose()

    failed = [w for ok, w in results if not ok]
    print(f"\n{len(results) - len(failed)}/{len(results)} checks passed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
