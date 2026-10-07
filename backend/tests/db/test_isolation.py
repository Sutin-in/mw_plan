"""Wave 10A: department data isolation across every read route (spec §22.1, §39; D-38 C-4).

A world holds data of Mock ENT (its own toner plan and a confirmed PPR), of Mock Store (a
Central Pool toner and an Assigned Purchase of gloves for Mock ENT, both with confirmed PPRs)
and a Plan Amendment of ENT's item. A requester of Mock OR - a department with none of it -
then calls EVERY GET route of the API, each path parameter pointing at the other departments'
records, and every report (screen and Excel). Each answer must be a refusal (403/404) or
contain none of their data. The routes are read from the application itself, so a new route
fails this test until it is given a target here.

D-41 (2026-10-05): a requester reads only the PPRs they created (even of their own
department), and may check a PR of any department before making a PPR for it - the PR
pre-validation route is therefore the one deliberate exception of the sweep.

Reference data (HOSxP masters) and organisation-level budget lines are not department data.
The user interface reads only through this API with the user's own token, so the API is where
isolation is enforced (UI hiding is not a security control, §39).
"""

from __future__ import annotations

import io
import re
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import Engine

from api_harness import FY, Harness, Plan
from ppr.domain.roles import Role

ENT, OR, STORE = "MOCK-DEP-ENT", "MOCK-DEP-OR", "MOCK-DEP-STORE"
TONER, PAPER, GLOVE = "MOCK-ITEM-TONER", "MOCK-ITEM-PAPER", "MOCK-ITEM-GLOVE"
OFFICE, MED = "MOCK-CAT-OFFICE", "MOCK-CAT-MED"
# Nothing of these may reach the OR requester.
MARKERS = ("PR-ENT-ONLY", "PR-STORE-C", "PR-STORE-A", "Mock toner", "Mock gloves", TONER, GLOVE)
# Not department data: HOSxP reference lists.
REFERENCE = {"/api/masters/{kind}"}
# D-41: any requester checks any department's PR (the PPR takes the PR's department).
ANY_PR = "/api/prs/{pr_no}/prevalidation"


class World:
    def __init__(self, hx: Harness) -> None:
        self.hx = hx
        self.plan = p = Plan(hx)
        p.sync()
        p.year()
        office = p.budget(OFFICE, "1600").json()["id"]
        med = p.budget(MED, "300").json()["id"]
        r = p.item(office, "1000")  # ENT toner 10 @ 100
        self.ent_item = int(r.json()["id"])
        assert p.item(office, "100", item_id=PAPER, owner_department_id=OR).status_code == 201
        r = p.item(
            office,
            "500",
            plan_type="CENTRAL",
            owner_department_id=STORE,
            purchasing_department_id=STORE,
            planned_qty="5",
        )
        self.central = int(r.json()["id"])
        r = p.item(
            med,
            "300",
            item_id=GLOVE,
            plan_type="ASSIGNED",
            owner_department_id=STORE,
            purchasing_department_id=STORE,
            demands=[{"department_id": ENT, "demand_qty": "10"}],
        )
        assert r.status_code == 201, r.text
        self.assigned = int(r.json()["id"])
        assert p.move("APPROVED").status_code == 200
        assert p.move("ACTIVE").status_code == 200
        self.ent = hx.h(hx.user("req_ent", Role.REQUESTER, department=ENT))
        self.store = hx.h(hx.user("req_store", Role.REQUESTER, department=STORE))
        self.observer = hx.h(hx.user("req_or", Role.REQUESTER, department=OR))
        self.ent_ppr = self.confirmed("PR-ENT-ONLY", ENT, self.ent, (TONER, "2", "100"))
        self.store_ppr = self.confirmed("PR-STORE-C", STORE, self.store, (TONER, "1", "100"))
        self.confirmed("PR-STORE-A", STORE, self.store, (GLOVE, "10", "30"))
        r = hx.client.post(
            f"/api/plan-years/{FY}/amendments",
            json={
                "approval_document_no": "D-ISO",
                "approval_date": "2026-10-14",
                "reason": "r",
                "items": [{"plan_item_id": self.ent_item, "estimated_unit_price": "90"}],
            },
            headers=p.h,
        )
        assert r.status_code == 201, r.text
        self.amendment = int(r.json()["id"])
        r = hx.client.post("/api/pr-sync", json={}, headers=p.h)
        assert r.status_code == 200, r.text
        self.run = int(r.json()["run"]["id"])

    def confirmed(
        self, pr_no: str, dept: str, h: dict[str, str], *lines: tuple[str, str, str]
    ) -> int:
        self.hx.put_pr(pr_no, *lines, dept=dept)
        r = self.hx.create_ppr(pr_no, h)
        assert r.status_code == 201, r.text
        c = self.hx.client.post(f"/api/pprs/{r.json()['id']}/confirm", headers=h)
        assert c.status_code == 200, c.text
        return int(r.json()["id"])


@pytest.fixture
def world(db: Engine) -> World:
    return World(Harness(db))


def _targets(w: World, template: str) -> list[str]:
    """The concrete paths that point the template at other departments' records."""
    values: dict[str, list[str]] = {
        "ppr_id": [str(w.ent_ppr), str(w.store_ppr)],
        "item_id": [str(w.ent_item), str(w.central), str(w.assigned)],
        "fiscal_year": [str(FY)],
        "amendment_id": [str(w.amendment)],
        "pr_no": ["PR-ENT-ONLY", "PR-STORE-C"],
        "run_id": [str(w.run)],
        "code": [x["code"] for x in w.hx.client.get("/api/reports", headers=w.plan.h).json()],
    }
    names = re.findall(r"\{([^}]+)\}", template)
    for n in names:
        assert n in values, f"give path parameter {{{n}}} of {template} a target in this test"
    paths = [template]
    for n in names:
        paths = [p.replace(f"{{{n}}}", v) for p in paths for v in values[n]]
    return paths


def _text(r: Any) -> str:
    if r.headers.get("content-type", "").startswith(
        "application/vnd.openxmlformats-officedocument"
    ):
        wb = load_workbook(io.BytesIO(r.content), read_only=True)
        return " ".join(
            str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row
        )
    return r.text


@pytest.mark.spec("S-22.1", "S-39", "AT-40.10.1", "D-38")
def test_a_requester_never_reads_another_departments_data(world: World) -> None:
    w = world
    templates = sorted(
        r.path  # type: ignore[attr-defined]
        for r in w.hx.client.app.routes  # type: ignore[attr-defined]
        if getattr(r, "path", "").startswith("/api/")
        and "GET" in getattr(r, "methods", set())
        and r.path not in REFERENCE  # type: ignore[attr-defined]
        and r.path != ANY_PR  # type: ignore[attr-defined]
    )
    checked: list[tuple[str, int]] = []
    leaks: list[str] = []
    for template in templates:
        for path in _targets(w, template):
            params = {"fiscal_year": FY} if path.startswith("/api/reports/") else None
            r = w.hx.client.get(path, params=params, headers=w.observer)
            checked.append((path, r.status_code))
            if r.status_code in (401, 403, 404, 409, 422):
                continue
            assert r.status_code == 200, (path, r.status_code, r.text[:200])
            body = _text(r)
            leaks += [f"{path}: {m}" for m in MARKERS if m in body]
    assert not leaks, leaks
    # The other departments' own records are refused outright, not just emptied.
    refused = dict(checked)
    for path in (
        f"/api/pprs/{w.ent_ppr}",
        f"/api/pprs/{w.store_ppr}/coverage",
        f"/api/plan-items/{w.ent_item}",
        f"/api/plan-items/{w.central}/balance",
        f"/api/plan-amendments/{w.amendment}",
    ):
        assert refused[path] == 404, path
    assert len(checked) > 60
    # D-41: the PR itself may be checked by any requester (no PPR of it is visible).
    r = w.hx.client.get("/api/prs/PR-ENT-ONLY/prevalidation", headers=w.observer)
    assert r.status_code == 200


@pytest.mark.spec("D-41", "S-22.1", "S-39", "AT-40.15.4")
def test_a_requester_never_reads_a_ppr_someone_else_created(world: World) -> None:
    """Same department, not the creator: none of the PPR routes or lists show it."""
    w = world
    colleague = w.hx.h(w.hx.user("req_ent_b", Role.REQUESTER, department=ENT))
    pid = w.ent_ppr
    # Every GET route on one PPR, read from the application itself (a new one fails here).
    ppr_routes = sorted(
        r.path.replace("{ppr_id}", str(pid))  # type: ignore[attr-defined]
        for r in w.hx.client.app.routes  # type: ignore[attr-defined]
        if getattr(r, "path", "").startswith("/api/pprs/{ppr_id}")
        and "GET" in getattr(r, "methods", set())
    )
    assert len(ppr_routes) >= 7, ppr_routes
    for path in ppr_routes:
        assert w.hx.client.get(path, headers=colleague).status_code == 404, path
    for path, params in (
        ("/api/pprs", None),
        ("/api/ppr-search", {"pr_no": "PR-ENT"}),
        ("/api/dashboard", {"fiscal_year": FY}),
        ("/api/alerts", None),
    ):
        r = w.hx.client.get(path, params=params, headers=colleague)
        assert r.status_code == 200 and "PR-ENT-ONLY" not in r.text, path
    dash = w.hx.client.get("/api/dashboard", params={"fiscal_year": FY}, headers=colleague)
    assert sum(dash.json()["ppr_counts"].values()) == 0 and dash.json()["recent"] == []
    assert dash.json()["own_pprs_only"] is True
    for x in w.hx.client.get("/api/reports", headers=w.plan.h).json():
        for suffix in ("", "/xlsx"):
            r = w.hx.client.get(
                f"/api/reports/{x['code']}{suffix}", params={"fiscal_year": FY}, headers=colleague
            )
            if x["code"] == "AUDIT":
                assert r.status_code == 403, x["code"]
                continue
            assert r.status_code == 200, (x["code"], suffix, r.text[:200])
            body = _text(r)
            assert "PR-ENT-ONLY" not in body and f"PPR-{FY}-" not in body, (x["code"], suffix)
    # Plan data stays by department: the colleague sees Mock ENT's plan item.
    assert "Mock toner" in w.hx.client.get(f"/api/plan-years/{FY}/items", headers=colleague).text


@pytest.mark.spec("S-22.1", "S-39", "S-37")
def test_a_department_filter_never_widens_a_requesters_scope(world: World) -> None:
    w = world
    for x in w.hx.client.get("/api/reports", headers=w.plan.h).json():
        r = w.hx.client.get(
            f"/api/reports/{x['code']}",
            params={"fiscal_year": FY, "department_id": ENT},
            headers=w.observer,
        )
        if r.status_code == 200:
            assert not [m for m in MARKERS if m in r.text], x["code"]
    r = w.hx.client.get("/api/ppr-search", params={"department_id": ENT}, headers=w.observer)
    assert r.status_code in (200, 403) and "PR-ENT-ONLY" not in r.text


@pytest.mark.spec("S-22.1", "D-38")
def test_the_owner_and_the_purchaser_do_see_their_own_records(world: World) -> None:
    # The same sweep would pass trivially if everything were refused: the departments
    # concerned do read their records.
    w = world
    ent_view = w.hx.client.get(f"/api/pprs/{w.ent_ppr}", headers=w.ent)  # its creator
    assert ent_view.status_code == 200 and "PR-ENT-ONLY" in ent_view.text
    store_items = w.hx.client.get(f"/api/plan-years/{FY}/items", headers=w.store).text
    assert "Mock toner" in store_items and "Mock gloves" in store_items
    ent_items = w.hx.client.get(f"/api/plan-years/{FY}/items", headers=w.ent).json()
    assert {i["id"] for i in ent_items} == {w.ent_item, w.assigned}  # its demand, not CENTRAL
