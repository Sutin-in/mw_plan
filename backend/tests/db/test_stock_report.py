"""Wave 10B-1 (D-39): central-store issues and returns in the Central Plan Usage report.

World: FY2570 ACTIVE (server date 2026-10-15), fund 1 OFFICE line 1500: Mock ENT paper
10 / 1000 (DEPARTMENT) and a CENTRAL toner 5 / 500 bought by Mock Store (toner's only plan
line); the store has confirmed a PPR for 4. MOCK HOSxP movements of toner: ENT issue 2, OR
issue 1, ENT return 1, and one before the fiscal year; and one paper movement of a kind the
site has not mapped (paper is not a Central Pool item, so it only shows in the note).
Everything here is made up; no HOSxP table, field or movement kind is implied.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from types import MappingProxyType
from typing import Any

import pytest
from sqlalchemy import Engine, text

from api_harness import Harness, Plan
from ppr.application.sync_services import NIGHTLY, run_sync
from ppr.domain.pr_sync import SyncScope
from ppr.domain.roles import Role
from ppr.infrastructure.db.repositories import unit_of_work
from ppr.integration.hosxp.contracts import (
    StockMovement,
    StockMovementKind,
    StockMovementSource,
    StockVerification,
)
from ppr.integration.hosxp.errors import HosxpConfigurationError

TONER, PAPER = "MOCK-ITEM-TONER", "MOCK-ITEM-PAPER"
OFFICE = "MOCK-CAT-OFFICE"
ENT, OR, STORE = "MOCK-DEP-ENT", "MOCK-DEP-OR", "MOCK-DEP-STORE"
NO_DATA = "ยังไม่มีข้อมูล"
URL = "/api/reports/CENTRAL_PLAN_USAGE"

VERIFIED = StockMovementSource(
    problem=None,
    verification=StockVerification("MOCK IT", date(2026, 10, 1), "MOCK-REF-10B"),
    kind_map=MappingProxyType(
        {"MOCK-ISSUE": StockMovementKind.ISSUE, "MOCK-RETURN": StockMovementKind.RETURN}
    ),
)


def move(
    i: int, day: int, dept: str, qty: str, kind: str, month: int = 10, item: str = TONER
) -> StockMovement:
    return StockMovement(
        movement_id=f"MOCK-SM-{i}",
        movement_date=date(2026, month, day),
        department_id=dept,
        item_id=item,
        qty=Decimal(qty),
        native_kind=kind,
    )


MOVES = [
    move(1, 2, ENT, "2", "MOCK-ISSUE"),
    move(2, 3, OR, "1", "MOCK-ISSUE"),
    move(3, 4, ENT, "1", "MOCK-RETURN"),
    move(4, 5, ENT, "5", "MOCK-ADJ", item=PAPER),  # unmapped kind, on a non-central item
    move(5, 30, ENT, "9", "MOCK-ISSUE", month=9),  # before the fiscal year: outside the range
]


class World:
    def __init__(
        self,
        hx: Harness,
        *,
        second_fund: bool = False,
        buyer2: str = STORE,
        ent_toner: bool = False,
    ) -> None:
        self.hx = hx
        self.plan = p = Plan(hx)
        p.sync()
        p.year()
        office = p.budget(OFFICE, "1500").json()["id"]
        # ENT's own DEPARTMENT line: paper, or (ent_toner) toner - then toner is also bought
        # outside the Central Plan and its HOSxP issues cannot be tied to the central line.
        assert p.item(office, "1000", item_id=TONER if ent_toner else PAPER).status_code == 201
        r = p.item(
            office,
            "500",
            plan_type="CENTRAL",
            owner_department_id=ENT,
            purchasing_department_id=STORE,
            planned_qty="5",
        )
        assert r.status_code == 201, r.text
        self.central_id = int(r.json()["id"])
        if second_fund:  # the same HOSxP item in a second CENTRAL Plan Item, another fund
            office2 = p.budget(OFFICE, "400", fund="MOCK-FUND-2").json()["id"]
            r = p.item(
                office2,
                "400",
                plan_type="CENTRAL",
                owner_department_id=buyer2,
                purchasing_department_id=buyer2,
                planned_qty="4",
            )
            assert r.status_code == 201, r.text
        assert p.move("APPROVED").status_code == 200
        assert p.move("ACTIVE").status_code == 200
        self.store = hx.h(hx.user("req_store", Role.REQUESTER, department=STORE))
        self.ent = hx.h(hx.user("req_ent", Role.REQUESTER, department=ENT))
        self.or_ = hx.h(hx.user("req_or", Role.REQUESTER, department=OR))
        self.ward = hx.h(hx.user("req_ward", Role.REQUESTER, department="MOCK-DEP-WARD-A"))
        hx.put_pr("PR-S1", (TONER, "4", "100"), dept=STORE)
        d = hx.create_ppr("PR-S1", self.store)
        assert d.status_code == 201, d.text
        ok = hx.client.post(f"/api/pprs/{d.json()['id']}/confirm", headers=self.store)
        assert ok.status_code == 200, ok.text
        self.gw = hx.gateway
        self.gw.data.stock_movements = list(MOVES)
        self.gw.data.stock_source = VERIFIED

    def report(self, h: dict[str, str], **params: str) -> dict[str, Any]:
        got = self.hx.client.get(URL, headers=h, params=params)
        assert got.status_code == 200, got.text
        return dict(got.json())

    def rows(self, h: dict[str, str], **params: str) -> list[dict[str, Any]]:
        return list(self.report(h, **params)["rows"])


@pytest.fixture
def w(db: Engine) -> World:
    return World(Harness(db))


def num(v: Any) -> Any:
    return Decimal(v) if isinstance(v, str) and v != NO_DATA else v


def stock(row: dict[str, Any]) -> tuple[Any, ...]:
    keys = ("issued_qty", "returned_qty", "net_used_qty", "not_issued_qty")
    return tuple(num(row.get(k)) for k in keys)


def xlsx_text(w: World, h: dict[str, str]) -> str:
    """Every cell of the downloaded Excel file, as one text."""
    from io import BytesIO

    from openpyxl import load_workbook

    got = w.hx.client.get(f"{URL}/xlsx", headers=h)
    assert got.status_code == 200 and got.content[:2] == b"PK"
    ws = load_workbook(BytesIO(got.content)).active
    return " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


def plan_data(engine: Engine) -> dict[str, list[tuple[Any, ...]]]:
    """Every plan and PPR table, row by row (the data D-39 says stock never changes).
    ``ppr_sync_observation`` is left out: it is the synchronization's own log (one row per PR
    looked at in a run), not plan or PPR data."""
    with engine.connect() as c:
        tables = [
            r[0]
            for r in c.execute(
                text(
                    "SELECT table_name FROM information_schema.tables WHERE table_schema = "
                    "current_schema() AND (table_name LIKE 'plan%' OR table_name LIKE 'ppr%') "
                    "AND table_name <> 'ppr_sync_observation' "
                    "ORDER BY table_name"
                )
            )
        ]
        return {
            t: [tuple(r) for r in c.execute(text(f"SELECT * FROM {t} ORDER BY 1"))] for t in tables
        }


# ------------------------------------------------------------------ AT-40.12.1
@pytest.mark.spec("AT-40.12.1", "D-39", "S-37")
def test_verified_movements_are_shown_per_item_and_department(w: World) -> None:
    for h in (w.plan.h, w.store):  # organisation-wide and the purchasing department
        rep = w.report(h)
        item, ent, or_ = rep["rows"]
        assert (item["row_kind"], item["purchasing_department"], item["fund_source"]) == (
            "รายการแผน",
            "Mock Store",
            "Mock fund 1",
        )
        assert (num(item["planned_qty"]), num(item["purchased_qty"])) == (5, 4)
        assert (num(item["remaining_qty"]), num(item["remaining_amount"])) == (1, 100)
        # issued 2 + 1, returned 1, net used 2, bought 4 - 2 = 2 not yet issued
        assert stock(item) == (3, 1, 2, 2)
        assert item["note"] == ""
        assert (ent["row_kind"], ent["department"], ent["item_code"]) == (
            "หน่วยงานเบิก/คืน",
            "Mock ENT",
            "I1",
        )
        assert stock(ent) == (2, 1, 1, None) and ent.get("planned_qty") is None
        assert (or_["department"], stock(or_)) == ("Mock OR", (1, 0, 1, None))
        notes = " ".join(rep["notes"])
        assert "MOCK IT" in notes and "MOCK-REF-10B" in notes and "2026-10-01" in notes
        assert "ไม่ลดหรือเพิ่มแผน" in notes
        assert "ข้อมูลไม่ครบ" in notes and "1 รายการ" in notes  # the unmapped adjustment
    xlsx = w.hx.client.get(f"{URL}/xlsx", headers=w.store)
    assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"


@pytest.mark.spec("D-39")
def test_issuing_more_than_bought_is_shown_not_hidden(w: World) -> None:
    w.gw.data.stock_movements.append(move(6, 6, OR, "4", "MOCK-ISSUE"))
    item = w.rows(w.plan.h)[0]
    assert stock(item) == (7, 1, 6, -2)  # bought 4, net used 6
    assert "เบิกมากกว่าที่ซื้อ" in item["note"]


@pytest.mark.spec("D-39")
def test_without_a_return_kind_no_net_figure_is_guessed(w: World) -> None:
    w.gw.data.stock_source = StockMovementSource(
        problem=None,
        verification=VERIFIED.verification,
        kind_map=MappingProxyType({"MOCK-ISSUE": StockMovementKind.ISSUE}),
    )
    # (ENT's "MOCK-RETURN" would now be an unmapped kind - see the next test; leave it out
    # here to look at the missing RETURN mapping alone.)
    w.gw.data.stock_movements = [x for x in MOVES if x.native_kind != "MOCK-RETURN"]
    rep = w.report(w.plan.h)
    item, ent, _ = rep["rows"]
    assert stock(item) == (3, NO_DATA, NO_DATA, NO_DATA)
    assert stock(ent) == (2, NO_DATA, NO_DATA, None)
    assert any("ประเภทการรับคืน" in n for n in rep["notes"])


# ------------------------------------------------------------------ visibility (D-39, C-4)
@pytest.mark.spec("D-39", "S-22.1", "S-39")
def test_a_department_sees_only_its_own_consumption(w: World) -> None:
    (ent,) = w.rows(w.ent)
    assert (ent["row_kind"], ent["department"], ent["item_name"]) == (
        "หน่วยงานเบิก/คืน",
        "Mock ENT",
        "Mock toner",
    )
    assert stock(ent) == (2, 1, 1, None)
    for hidden in (
        "planned_qty",
        "purchased_qty",
        "remaining_qty",
        "planned_amount",
        "purchased_amount",
        "remaining_amount",
        "fund_source",
        "budget_category",
        "purchasing_department",
    ):
        assert ent.get(hidden) in (None, ""), hidden
    (or_,) = w.rows(w.or_)
    assert (or_["department"], stock(or_)) == ("Mock OR", (1, 0, 1, None))
    assert w.rows(w.ward) == []  # received nothing: sees nothing
    # A filter never widens the scope, nor probes plan details the department cannot read.
    assert w.rows(w.ent, department_id=OR) == []
    assert w.rows(w.ent, department_id=STORE) == []
    assert w.rows(w.ent, fund_source_id="MOCK-FUND-1") == []
    assert w.rows(w.ent, budget_category_id=OFFICE) == []
    assert len(w.rows(w.ent, department_id=ENT)) == 1
    assert len(w.rows(w.ent, item="toner")) == 1 and w.rows(w.ent, item="paper") == []
    text_all = str(w.report(w.ent))
    assert "Mock OR" not in text_all and "Mock Store" not in text_all
    # The hospital-wide count of unmapped movements is not given to a single department.
    assert any("บางรายการ" in n for n in w.report(w.ent)["notes"])
    # The Excel file holds exactly what the screen shows.
    sheet = xlsx_text(w, w.ent)
    assert "Mock ENT" in sheet and "Mock toner" in sheet
    for other in ("Mock OR", "Mock Store", "Mock fund", "Office", "รายการแผน"):
        assert other not in sheet, other
    sheet = xlsx_text(w, w.ward)
    assert "Mock toner" not in sheet and "Mock ENT" not in sheet and "Mock OR" not in sheet


@pytest.mark.spec("D-39", "S-22.1")
def test_a_department_sees_nothing_while_there_is_no_data(w: World) -> None:
    w.gw.data.stock_source = StockMovementSource()
    rep = w.report(w.ent)
    assert rep["rows"] == [] and any(NO_DATA in n for n in rep["notes"])


# ------------------------------------------------------------------ no fund split (D-39)
@pytest.mark.spec("D-39")
def test_stock_of_an_item_in_two_funds_is_shown_once_never_split(db: Engine) -> None:
    w = World(Harness(db), second_fund=True)
    rows = w.rows(w.plan.h)
    kinds = [r["row_kind"] for r in rows]
    assert kinds == [
        "รายการแผน",
        "รายการแผน",
        "รวมสินค้า (ทุกรายการแผน)",
        "หน่วยงานเบิก/คืน",
        "หน่วยงานเบิก/คืน",
    ]
    f1, f2, total = rows[:3]
    assert {f1["fund_source"], f2["fund_source"]} == {"Mock fund 1", "Mock fund 2"}
    for plan_row in (f1, f2):  # nothing of the stock is put on a fund's line
        assert stock(plan_row) == (None, None, None, None)
    assert (num(total["planned_qty"]), num(total["purchased_qty"])) == (9, 4)
    assert stock(total) == (3, 1, 2, 2)
    assert total["note"] == "ไม่สามารถแบ่งยอดตามแหล่งเงินได้จากข้อมูล HOSxP"
    # A fund filter shows that fund's plan line, and the item's stock still only as a whole.
    rows = w.rows(w.plan.h, fund_source_id="MOCK-FUND-2")
    assert [r["row_kind"] for r in rows][:2] == ["รายการแผน", "รวมสินค้า (ทุกรายการแผน)"]
    assert rows[0]["fund_source"] == "Mock fund 2" and stock(rows[0])[0] is None
    assert stock(rows[1]) == (3, 1, 2, 2) and "แหล่งเงิน" in rows[1]["note"]
    with w.hx.engine.connect() as c:  # and never posted against any fund
        assert (
            c.execute(
                text("SELECT count(*) FROM plan_ledger WHERE event_type LIKE 'CENTRAL_STOCK%'")
            ).scalar_one()
            == 0
        )


@pytest.mark.spec("D-39", "S-22.1")
def test_two_purchasers_of_one_item_never_mix_their_purchases(db: Engine) -> None:
    """Toner is bought by Mock Store (fund 1) and by Mock IT (fund 2). HOSxP issues cannot be
    told apart by purchaser: each purchaser sees the item's issues once, marked, and no
    "bought but not issued" made from its own purchases against everyone's issues."""
    w = World(Harness(db), second_fund=True, buyer2="MOCK-DEP-IT")
    rows = w.rows(w.store)
    plan, total = rows[0], rows[1]
    assert plan["row_kind"] == "รายการแผน" and stock(plan) == (None, None, None, None)
    assert total["row_kind"] == "รวมสินค้า (ทุกรายการแผน)"
    assert (num(total["planned_qty"]), num(total["purchased_qty"])) == (5, 4)  # Store's only
    assert stock(total) == (3, 1, 2, NO_DATA)
    assert "หน่วยงานผู้ซื้อ" in total["note"] and "เบิกมากกว่า" not in total["note"]
    assert "Mock IT" not in str(rows)  # the other purchaser's line stays hidden (C-4)
    # Organisation-wide, both lines are in view: the whole item can be compared.
    total = next(r for r in w.rows(w.plan.h) if r["row_kind"] == "รวมสินค้า (ทุกรายการแผน)")
    assert (num(total["planned_qty"]), num(total["purchased_qty"])) == (9, 4)
    assert stock(total) == (3, 1, 2, 2) and "และตามหน่วยงานผู้ซื้อ" in total["note"]
    # ... but not once narrowed to one purchaser.
    rows = w.rows(w.plan.h, department_id=STORE)
    total = next(r for r in rows if r["row_kind"] == "รวมสินค้า (ทุกรายการแผน)")
    assert stock(total)[3] == NO_DATA


# ------------------------------------------------------------------ AT-40.12.2 / .5
@pytest.mark.spec("AT-40.12.2", "D-39", "S-18.3")
def test_movements_never_change_the_plan_balance(w: World) -> None:
    before = plan_data(w.hx.engine)
    bal0 = w.hx.client.get(f"/api/plan-items/{w.central_id}/balance", headers=w.plan.h).json()
    w.rows(w.plan.h)
    w.gw.data.stock_movements.append(move(7, 7, ENT, "3", "MOCK-ISSUE"))
    w.rows(w.plan.h)
    bal1 = w.hx.client.get(f"/api/plan-items/{w.central_id}/balance", headers=w.plan.h).json()
    assert bal0 == bal1
    assert (bal1["remaining_qty"], bal1["remaining_amount"]) == ("1.0000", "100.00")
    assert not any("stock" in k for k in bal1)
    assert plan_data(w.hx.engine) == before


@pytest.mark.spec("AT-40.12.5", "D-39", "S-26")
@pytest.mark.parametrize(
    ("case", "reason"),
    [
        ("not_configured", "ยังไม่ได้ตั้งค่า"),
        ("not_verified", "ยังไม่ได้รับการยืนยัน"),
        ("hosxp_down", "เชื่อมต่อ HOSxP ไม่ได้"),
        ("duplicate_id", "ไม่ตรงตามที่กำหนด"),
        ("query_error", "ทำงานไม่สำเร็จ"),
    ],
)
def test_no_verified_data_shows_no_data_and_touches_nothing(
    w: World, case: str, reason: str
) -> None:
    before = plan_data(w.hx.engine)
    if case == "not_configured":
        w.gw.data.stock_source = StockMovementSource()
    elif case == "not_verified":
        w.gw.data.stock_source = StockMovementSource(
            problem="query ยอดเบิก/คืนยังไม่ได้รับการยืนยัน (ไม่มี verified_by / reference)"
        )
    elif case == "hosxp_down":
        w.gw.available = False
    elif case == "duplicate_id":
        w.gw.data.stock_movements.append(MOVES[0])
    else:

        def broken(*_: Any) -> Any:
            raise HosxpConfigurationError("get_stock_movements: missing column(s) ['qty']")

        w.gw.get_stock_movements = broken
    rep = w.report(w.plan.h)
    item, *rest = rep["rows"]
    assert rest == []  # no department rows without data
    assert stock(item) == (NO_DATA, NO_DATA, NO_DATA, NO_DATA)
    assert (num(item["purchased_qty"]), num(item["remaining_amount"])) == (4, 100)  # PPR part
    assert any(NO_DATA in n and reason in n for n in rep["notes"]), rep["notes"]
    assert w.hx.client.get(f"{URL}/xlsx", headers=w.plan.h).status_code == 200
    w.gw.available = True
    # The rest of the system is unaffected: a PPR can still be drafted against the plan.
    w.hx.put_pr("PR-S2", (TONER, "1", "100"), dept=STORE)
    assert w.hx.create_ppr("PR-S2", w.store).status_code == 201
    after = plan_data(w.hx.engine)
    assert after["plan_ledger"] == before["plan_ledger"]
    assert after["plan_item"] == before["plan_item"]


# ------------------------------------------------------------------ AT-40.12.6
@pytest.mark.spec("AT-40.12.6", "D-39", "S-25.5")
def test_repeating_reports_reads_and_synchronization_changes_no_plan_or_ppr_data(
    w: World,
) -> None:
    before = plan_data(w.hx.engine)
    first = w.rows(w.plan.h)
    for _ in range(3):
        assert w.rows(w.plan.h) == first
        assert w.rows(w.store) == first
        w.rows(w.ent)
        assert w.hx.client.get(f"{URL}/xlsx", headers=w.plan.h).status_code == 200
        r = run_sync(
            lambda: unit_of_work(w.hx.engine),
            w.gw,
            mode=NIGHTLY,
            scope=SyncScope.PPRS,
            requested_by=None,
        )
        assert r.run.status == "SUCCEEDED"
    assert plan_data(w.hx.engine) == before
    assert w.gw.calls.count("get_stock_movements") == 1 + 3 * 4  # read live every time


# ------------------------------------------------------------------ AT-40.12.7
@pytest.mark.spec("AT-40.12.7", "D-39", "S-27")
def test_hosxp_stays_the_source_of_truth_for_stock(w: World) -> None:
    with w.hx.engine.connect() as c:
        found = c.execute(
            text(
                "SELECT table_name, column_name FROM information_schema.columns "
                "WHERE table_schema = current_schema() AND (table_name ~* 'stock|movement' "
                "OR column_name ~* 'stock|movement')"
            )
        ).all()
    assert found == []  # no stock ledger, no stored movements
    assert stock(w.rows(w.plan.h)[0]) == (3, 1, 2, 2)
    # HOSxP changes (a movement corrected, another added): the next report shows it as is.
    w.gw.data.stock_movements[0] = move(1, 2, ENT, "1", "MOCK-ISSUE")
    w.gw.data.stock_movements.append(move(8, 8, OR, "1", "MOCK-RETURN"))
    item, ent, or_ = w.rows(w.plan.h)
    assert stock(item) == (2, 2, 0, 4)
    assert stock(ent) == (1, 1, 0, None) and stock(or_) == (1, 1, 0, None)
    w.gw.data.stock_movements.clear()  # gone from HOSxP: gone from the report
    assert w.rows(w.plan.h)[1:] == [] and stock(w.rows(w.plan.h)[0]) == (0, 0, 0, 4)


# ------------------------------------------------------------------ unknown / ambiguous
@pytest.mark.spec("AT-40.12.5", "D-39")
def test_an_unmapped_movement_makes_its_figures_unknown_never_zero(w: World) -> None:
    """Unknown is never coerced to zero: OR's adjustment could be an issue or a return."""
    w.gw.data.stock_movements.append(move(9, 9, OR, "4", "MOCK-ADJ"))
    rows = w.rows(w.plan.h)
    item = rows[0]
    by_dept = {r["department"]: r for r in rows[1:]}
    assert stock(item) == (NO_DATA, NO_DATA, NO_DATA, NO_DATA)  # the item total is unknown
    assert stock(by_dept["Mock OR"]) == (NO_DATA, NO_DATA, NO_DATA, None)
    assert stock(by_dept["Mock ENT"]) == (2, 1, 1, None)  # ENT's own figures are complete
    assert any("ไม่นับเป็นศูนย์" in n for n in w.report(w.plan.h)["notes"])
    (own,) = w.rows(w.or_)  # OR itself sees its row as unknown, not as 1
    assert stock(own) == (NO_DATA, NO_DATA, NO_DATA, None)
    # A department with nothing but an unmapped movement still gets an (unknown) row.
    w.gw.data.stock_movements.append(move(10, 10, "MOCK-DEP-WARD-A", "1", "MOCK-ADJ"))
    (ward,) = w.rows(w.ward)
    assert stock(ward) == (NO_DATA, NO_DATA, NO_DATA, None)


@pytest.mark.spec("D-39")
def test_issues_of_an_item_also_bought_outside_the_central_plan_are_not_attributed(
    db: Engine,
) -> None:
    """Toner is also on ENT's own DEPARTMENT plan: HOSxP cannot say which purchase an issue
    came from, so the issues are shown for the item, never put on the central line, and
    "bought but not issued" is not computed."""
    w = World(Harness(db), ent_toner=True)
    for h in (w.plan.h, w.store):
        rows = w.rows(h)
        plan, total = rows[0], rows[1]
        assert plan["row_kind"] == "รายการแผน" and stock(plan) == (None, None, None, None)
        assert total["row_kind"] == "รวมสินค้า (ทุกรายการแผน)"
        assert (num(total["planned_qty"]), num(total["purchased_qty"])) == (5, 4)  # central
        assert stock(total) == (3, 1, 2, NO_DATA)
        assert "แผนประเภทอื่น" in total["note"]
        assert [r["department"] for r in rows[2:]] == ["Mock ENT", "Mock OR"]
