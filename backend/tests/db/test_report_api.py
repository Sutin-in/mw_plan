"""Wave 8B through the API: reports and Excel export (spec §37, §22.1, §36; D-35).

World (test_ppr_api): FY2570 ACTIVE, Mock ENT fund 1 - toner 10 / 1000, paper 10 / 500,
gloves 10 / 300; Mock OR fund 2 - paper 10 / 100. Server date 2026-10-15.
"""

from __future__ import annotations

from datetime import UTC, datetime, timedelta
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import Engine, text
from test_ppr_api import GLOVE, PAPER, TONER, World

from api_harness import FY, Harness
from ppr.domain.roles import Role


class Rep:
    def __init__(self, w: World) -> None:
        self.w = w
        self.hx = w.hx
        self.officer = w.plan.h
        self.admin = self.hx.h(self.hx.user("admin", Role.ADMIN))
        self.cfo = self.hx.h(self.hx.user("cfo", Role.CFO))
        self.or_req = self.hx.h(self.hx.user("req_or", Role.REQUESTER, department="MOCK-DEP-OR"))

    def get(self, code: str, headers: dict[str, str], **q: Any) -> Any:
        r = self.hx.client.get(f"/api/reports/{code}", params=q, headers=headers)
        assert r.status_code == 200, r.text
        return r.json()

    def rows(self, code: str, headers: dict[str, str], **q: Any) -> list[dict[str, Any]]:
        return list(self.get(code, headers, **q)["rows"])

    def confirmed(self, pr_no: str, *lines: tuple[str, str, str]) -> dict[str, Any]:
        d = self.w.draft(pr_no, *lines)
        r = self.w.confirm(d["id"])
        assert r.status_code == 200, r.text
        return dict(r.json())

    def set_state(self, ppr_id: int, state: str) -> None:
        with self.hx.engine.begin() as c:
            c.execute(text("UPDATE ppr SET state = :s WHERE id = :i"), {"s": state, "i": ppr_id})

    def count(self, table: str) -> int:
        with self.hx.engine.connect() as c:
            return int(c.execute(text(f"SELECT count(*) FROM {table}")).scalar_one())


@pytest.fixture
def r(db: Engine) -> Rep:
    return Rep(World(Harness(db)))


@pytest.mark.spec("D-35", "S-37")
def test_the_reports_are_listed_and_audit_needs_audit_read(r: Rep) -> None:
    listed = {
        x["code"]: x["available"] for x in r.hx.client.get("/api/reports", headers=r.cfo).json()
    }
    assert set(listed) == {
        "PLAN_SUMMARY",
        "PLAN_REMAINING",
        "DEPARTMENT_UTILIZATION",
        "CATEGORY_UTILIZATION",
        "PPR_REGISTER",
        "PENDING_PPR",
        "CANCELLED_PPR",
        "PR_CHANGED",
        "AUDIT",
        "PLAN_AMENDMENT",  # Wave 7A (D-37)
        "CENTRAL_PLAN_USAGE",  # Wave 7B-1 (D-38)
    }
    assert listed["AUDIT"] is False and listed["PPR_REGISTER"] is True
    got = r.hx.client.get("/api/reports/AUDIT", headers=r.cfo)
    assert got.status_code == 403
    assert r.hx.client.get("/api/reports/NOPE", headers=r.cfo).status_code == 404
    assert r.get("AUDIT", r.officer)["rows"]


@pytest.mark.spec("D-35", "S-37", "S-18.3")
def test_plan_reports_follow_the_ledger_and_default_to_the_current_year(r: Rep) -> None:
    r.confirmed("PR-R1", (TONER, "9", "100"), (PAPER, "2", "50"))
    summary = r.get("PLAN_SUMMARY", r.cfo)
    assert ("ปีงบประมาณ", str(FY)) in [tuple(a) for a in summary["applied"]]
    office = next(
        x
        for x in summary["rows"]
        if x["budget_category"] == "Office supplies" and x["fund_source"] == "Mock fund 1"
    )
    assert (office["approved_budget"], office["planned_amount"], office["used_amount"]) == (
        "1500.00",
        "1500.00",
        "1000.00",
    )
    assert office["used_percent"] == "66.67" and office["items"] == 2 and office["low"] == 1
    rem = {x["item_code"]: x for x in r.rows("PLAN_REMAINING", r.cfo, fiscal_year=FY)}
    toner = next(v for v in rem.values() if v["item_name"] == "Mock toner")
    assert (toner["remaining_qty"], toner["remaining_amount"], toner["alert"]) == (
        "1.0000",
        "100.00",
        "ใกล้หมด",
    )
    depts = {x["department"]: x["used_amount"] for x in r.rows("DEPARTMENT_UTILIZATION", r.cfo)}
    assert depts == {"Mock ENT": "1000.00", "Mock OR": "0.00"}
    cats = {x["budget_category"]: x["items"] for x in r.rows("CATEGORY_UTILIZATION", r.cfo)}
    assert cats["Office supplies"] == 2 and len(cats) == 3


@pytest.mark.spec("D-35", "S-22.1", "D-41", "AT-40.15.4")
def test_a_requester_gets_only_their_department_and_no_envelope(r: Rep) -> None:
    r.confirmed("PR-R2", (TONER, "1", "100"))
    rows = r.rows("PLAN_SUMMARY", r.or_req)
    assert [(x["fund_source"], x["approved_budget"], x["items"]) for x in rows] == [
        ("Mock fund 2", None, 1)
    ]
    assert [x["department"] for x in r.rows("PLAN_REMAINING", r.or_req)] == ["Mock OR"]
    assert r.rows("PPR_REGISTER", r.or_req) == []
    assert len(r.rows("PPR_REGISTER", r.w.req)) == 1
    # D-41: PPR reports of a requester are the PPRs they created; plan reports by department
    assert ["ขอบเขต", "เฉพาะ PPR ที่คุณสร้าง"] in r.get("PPR_REGISTER", r.or_req)["applied"]
    assert ["ขอบเขต", "เฉพาะหน่วยงาน Mock OR"] in r.get("PLAN_SUMMARY", r.or_req)["applied"]


@pytest.mark.spec("D-35", "S-37")
def test_ppr_reports_use_their_states_and_the_first_confirmation_date(r: Rep) -> None:
    r.confirmed("PR-A", (TONER, "1", "100"))
    b = r.confirmed("PR-B", (PAPER, "1", "50"))
    c = r.confirmed("PR-C", (GLOVE, "1", "30"))
    draft = r.w.draft("PR-D", (TONER, "1", "100"))
    r.set_state(b["id"], "CANCELLED")
    r.set_state(c["id"], "PR_CHANGED_REVIEW_REQUIRED")
    reg = [x["pr_no"] for x in r.rows("PPR_REGISTER", r.cfo)]
    assert sorted(reg) == ["PR-A", "PR-B", "PR-C"]  # no draft
    pending = sorted(x["pr_no"] for x in r.rows("PENDING_PPR", r.cfo))
    assert pending == ["PR-A", "PR-C", "PR-D"]
    assert [x["pr_no"] for x in r.rows("CANCELLED_PPR", r.cfo)] == ["PR-B"]
    assert [x["pr_no"] for x in r.rows("PR_CHANGED", r.cfo)] == ["PR-C"]
    # Date range = first confirmation (version 1); drafts have none (and the rows are
    # append-only, so the dates are those of today's confirmations).
    today = datetime.now(UTC).astimezone().date()
    same_day = r.get("PENDING_PPR", r.cfo, date_from=today.isoformat(), date_to=today.isoformat())
    assert sorted(x["pr_no"] for x in same_day["rows"]) == ["PR-A", "PR-C"]
    assert same_day["notes"]  # drafts excluded by the date range, said on the report
    before = (today - timedelta(days=1)).isoformat()
    assert r.rows("PPR_REGISTER", r.cfo, date_to=before) == []
    after = (today + timedelta(days=1)).isoformat()
    assert r.rows("PPR_REGISTER", r.cfo, date_from=after) == []
    one = r.rows("PPR_REGISTER", r.cfo, state="CANCELLED")
    assert [x["pr_no"] for x in one] == ["PR-B"]
    bad = r.hx.client.get("/api/reports/PENDING_PPR?state=CANCELLED", headers=r.cfo)
    assert bad.status_code == 422 and bad.json()["detail"]["code"] == "STATE_NOT_IN_REPORT"
    backwards = r.hx.client.get(
        "/api/reports/PPR_REGISTER?date_from=2026-10-10&date_to=2026-10-01", headers=r.cfo
    )
    assert backwards.status_code == 422
    assert draft["state"] == "DRAFT"


@pytest.mark.spec("D-35", "S-37", "S-36")
def test_excel_file_has_the_same_rows_numbers_as_numbers_and_is_audited(r: Rep) -> None:
    r.confirmed("PR-X", (TONER, "9", "100"))
    screen = r.get("PLAN_REMAINING", r.cfo)
    before = r.count("audit_log")
    ledger = r.count("plan_ledger")
    got = r.hx.client.get("/api/reports/PLAN_REMAINING/xlsx", headers=r.cfo)
    assert got.status_code == 200
    assert got.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="PLAN_REMAINING_' in got.headers["content-disposition"]
    ws = load_workbook(BytesIO(got.content)).active
    values = [[c.value for c in row] for row in ws.iter_rows()]
    header = next(i for i, row in enumerate(values) if row[0] == "รหัสสินค้า")
    data = [row for row in values[header + 1 :] if any(v is not None for v in row)]
    assert len(data) == len(screen["rows"])
    toner = next(row for row in data if row[1] == "Mock toner")
    assert toner[9] == 1 and toner[12] == 100 and isinstance(toner[12], int | float)
    assert r.count("audit_log") == before + 1 and r.count("plan_ledger") == ledger
    r.get("PLAN_REMAINING", r.cfo)  # viewing on screen writes nothing
    assert r.count("audit_log") == before + 1
    row = r.w.plan.audit("REPORT_EXPORTED")[-1]
    assert row[2]["rows"] == len(screen["rows"])


@pytest.mark.spec("D-35")
def test_text_that_looks_like_a_formula_stays_text(r: Rep) -> None:
    with r.hx.engine.begin() as c:
        c.execute(
            text(
                "INSERT INTO audit_log (actor_kind, action, entity_type, entity_id, reason) "
                "VALUES ('SYSTEM', 'X_TEST', 't', '1', '=HYPERLINK(\"http://x\")')"
            )
        )
    got = r.hx.client.get("/api/reports/AUDIT/xlsx?action=X_TEST", headers=r.admin)
    ws = load_workbook(BytesIO(got.content)).active
    cells = [c for row in ws.iter_rows() for c in row if c.value == '=HYPERLINK("http://x")']
    assert cells and all(c.data_type == "s" for c in cells)


# ------------------------------------------------------------------ review follow-ups
@pytest.mark.spec("D-35", "S-22.1")
def test_a_requester_cannot_widen_the_scope_or_take_the_audit_file(r: Rep) -> None:
    r.confirmed("PR-S1", (TONER, "1", "100"))
    assert r.rows("PPR_REGISTER", r.or_req, department_id="MOCK-DEP-ENT") == []
    before = r.count("audit_log")
    got = r.hx.client.get("/api/reports/AUDIT/xlsx", headers=r.w.req)
    assert got.status_code == 403 and r.count("audit_log") == before


@pytest.mark.spec("D-35")
def test_filters_a_report_does_not_list_are_ignored(r: Rep) -> None:
    full = r.rows("PLAN_SUMMARY", r.cfo)
    same = r.get("PLAN_SUMMARY", r.cfo, department_id="MOCK-DEP-OR", item="toner")
    assert same["rows"] == full
    assert all(a[0] not in ("หน่วยงาน", "รายการ") for a in same["applied"])


@pytest.mark.spec("D-35")
def test_long_reports_are_capped_and_say_so(r: Rep, monkeypatch: pytest.MonkeyPatch) -> None:
    from ppr.application import report_services

    r.confirmed("PR-M1", (TONER, "1", "100"))
    r.confirmed("PR-M2", (PAPER, "1", "50"))
    monkeypatch.setattr(report_services, "MAX_ROWS", 1)
    got = r.get("PPR_REGISTER", r.cfo)
    assert len(got["rows"]) == 1 and got["truncated"] is True and got["notes"]
    assert got["rows"][0]["pr_no"] == "PR-M2"  # newest first
    audit = r.get("AUDIT", r.officer)
    assert len(audit["rows"]) == 1 and audit["truncated"] is True


@pytest.mark.spec("D-35")
def test_control_characters_do_not_break_the_file_and_extreme_dates_are_refused(r: Rep) -> None:
    with r.hx.engine.begin() as c:
        c.execute(
            text(
                "INSERT INTO audit_log (actor_kind, action, entity_type, entity_id, reason) "
                "VALUES ('SYSTEM', 'X_CTRL', 't', '1', :reason)"
            ),
            {"reason": "bad\x01text\x0b!"},
        )
    got = r.hx.client.get("/api/reports/AUDIT/xlsx?action=X_CTRL", headers=r.admin)
    assert got.status_code == 200
    ws = load_workbook(BytesIO(got.content)).active
    assert "badtext!" in [c.value for row in ws.iter_rows() for c in row]
    far = r.hx.client.get("/api/reports/PPR_REGISTER?date_to=9999-12-31", headers=r.cfo)
    assert far.status_code == 422 and far.json()["detail"]["code"] == "INVALID_DATE_RANGE"
