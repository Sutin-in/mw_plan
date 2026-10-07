"""Wave 12A-1 (D-42): importing a plan from Excel - preview, all-or-nothing import, rows
without a HOSxP item, removal while DRAFT.

The workbook is built here with openpyxl in the template's layout (sheets ``หัวข้อ`` and
``รายการ``); the mock HOSxP masters are synchronized first.
"""

from __future__ import annotations

import hashlib
from collections.abc import Sequence
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import Engine, text

from api_harness import FY, Harness, Plan
from ppr.domain.plan_import import HCol, RCol
from ppr.domain.roles import Role

HEAD_TITLES = [c.value for c in HCol]
ROW_TITLES = [c.value for c in RCol]


def workbook(
    headings: Sequence[Sequence[Any]], rows: Sequence[Sequence[Any]], *, drop: str | None = None
) -> bytes:
    wb = Workbook()
    wb.active.title = "คำอธิบาย"
    for name, titles, data in (("หัวข้อ", HEAD_TITLES, headings), ("รายการ", ROW_TITLES, rows)):
        if name == drop:
            continue
        ws = wb.create_sheet(name)
        ws.append(titles)
        for r in data:
            ws.append(list(r))
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


# heading: code, name, plan type, owner, purchaser, fund, category, note
H_ENT = [
    "H1",
    "ENT office",
    "DEPARTMENT",
    "MOCK-DEP-ENT",
    None,
    "MOCK-FUND-1",
    "MOCK-CAT-OFFICE",
    None,
]
H_CEN = [
    "H2",
    "Store gloves",
    "CENTRAL",
    "MOCK-DEP-STORE",
    "MOCK-DEP-STORE",
    "MOCK-FUND-1",
    "MOCK-CAT-MED",
    None,
]


def row(
    heading: str, name: str, unit: str, qty: Any, price: Any, amount: Any, **kw: Any
) -> list[Any]:
    """heading, name, unit, qty, price, amount, q1..q4, source sheet, note (no item, D-43)"""
    return [
        heading,
        name,
        unit,
        qty,
        price,
        amount,
        kw.get("q1"),
        kw.get("q2"),
        kw.get("q3"),
        kw.get("q4"),
        kw.get("sheet"),
        kw.get("note"),
    ]


GOOD_ROWS = [
    row("H1", "กระดาษ A4 80 แกรม", "รีม", 100, 120, 12000, q1=6000, q2=6000, sheet="สำนักงาน"),
    row("H1", "หมึกพิมพ์", "box", 4, 1500, 6000),
    row("H1", "ปากกา (ไม่ซื้อปีนี้)", "ด้าม", 0, 10, 0),  # skipped
    row("H2", "ถุงมือ", "กล่อง", 50, 80.5, 4025),
    row("H2", "ถุงมือ ไซซ์พิเศษ", "กล่อง", 3, 99.99, 300),  # amount != qty x price: warning
]
GOOD = workbook([H_ENT, H_CEN], GOOD_ROWS)


@pytest.fixture
def hx(db: Engine) -> Harness:
    return Harness(db)


@pytest.fixture
def plan(hx: Harness) -> Plan:
    p = Plan(hx)
    p.sync()
    p.year()
    assert p.budget("MOCK-CAT-OFFICE", "18000").status_code == 201
    assert p.budget("MOCK-CAT-MED", "4325").status_code == 201
    return p


def post(p: Plan, path: str, data: bytes, **params: Any) -> Any:
    return p.hx.client.post(
        path,
        content=data,
        params=params,
        headers={**p.h, "Content-Type": "application/octet-stream"},
    )


def preview(p: Plan, data: bytes) -> Any:
    return post(p, f"/api/plan-years/{FY}/imports/preview", data, file_name="แผน.xlsx")


def do_import(p: Plan, data: bytes, **params: Any) -> Any:
    return post(p, f"/api/plan-years/{FY}/imports", data, file_name="แผน.xlsx", **params)


def count(engine: Engine, sql: str) -> int:
    with engine.connect() as conn:
        return int(conn.execute(text(sql)).scalar_one())


# ------------------------------------------------------------------ template
@pytest.mark.spec("D-42", "AT-40.17.1")
def test_the_template_lists_the_hosxp_codes_and_budget_lines(plan: Plan) -> None:
    r = plan.hx.client.get(f"/api/plan-years/{FY}/import-template", headers=plan.h)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = load_workbook(BytesIO(r.content))
    assert {
        "คำอธิบาย",
        "หัวข้อ",
        "รายการ",
        "หน่วยงาน HOSxP",
        "แหล่งเงิน HOSxP",
        "หมวดงบ HOSxP",
        "วงเงินในระบบ",
    } <= set(wb.sheetnames)
    assert [c.value for c in wb["หัวข้อ"][1]] == HEAD_TITLES
    assert [c.value for c in wb["รายการ"][1]] == ROW_TITLES
    depts = {r[0] for r in wb["หน่วยงาน HOSxP"].iter_rows(min_row=2, values_only=True)}
    assert "MOCK-DEP-ENT" in depts and "MOCK-DEP-OLD" not in depts  # active only
    budgets = list(wb["วงเงินในระบบ"].iter_rows(min_row=2, values_only=True))
    assert ("MOCK-FUND-1", "Mock fund 1", "MOCK-CAT-OFFICE", "Office supplies", 18000) in budgets
    # an empty template imports nothing
    empty = preview(plan, r.content).json()
    assert not empty["can_import"] and empty["errors"][0]["code"] == "NO_ROWS"


# ------------------------------------------------------------------ preview + import
@pytest.mark.spec("D-42", "AT-40.17.1", "AT-40.17.3")
def test_a_valid_file_is_previewed_then_imported_once(plan: Plan) -> None:
    hx = plan.hx
    pv = preview(plan, GOOD)
    assert pv.status_code == 200, pv.text
    body = pv.json()
    assert body["can_import"] and body["errors_total"] == 0
    assert (body["rows_to_import"], body["skipped_total"]) == (4, 1)
    assert "pre_bound" not in body  # D-43: no HOSxP item codes in the plan
    assert Decimal(body["total_amount"]) == Decimal("22325")
    assert body["skipped"][0]["code"] == "ZERO_ROW" and body["skipped"][0]["row"] == 4
    assert [w["code"] for w in body["warnings"]] == ["AMOUNT_DIFFERS_FROM_ESTIMATE"]
    lines = {(b["budget_category_id"], Decimal(b["difference"])) for b in body["budgets"]}
    assert lines == {("MOCK-CAT-OFFICE", Decimal(0)), ("MOCK-CAT-MED", Decimal(0))}
    assert body["file_sha256"] == hashlib.sha256(GOOD).hexdigest()
    assert count(hx.engine, "SELECT count(*) FROM plan_item") == 0  # preview writes nothing

    r = do_import(plan, GOOD, expected_sha256=body["file_sha256"])
    assert r.status_code == 201, r.text
    imp = r.json()
    assert (imp["rows_imported"], imp["rows_skipped"], imp["state"]) == (4, 1, "IMPORTED")
    items = hx.client.get(f"/api/plan-years/{FY}/items", headers=plan.h).json()
    assert len(items) == 4
    by_name = {i["item_name"]: i for i in items}
    paper = by_name["กระดาษ A4 80 แกรม"]
    assert (paper["item_id"], paper["item_code"], paper["unit"]) == (None, None, "รีม")
    assert (paper["plan_import_id"], paper["source_sheet"], paper["source_row"]) == (
        imp["id"],
        "สำนักงาน",
        2,
    )
    assert (Decimal(paper["q1"]), Decimal(paper["q2"]), paper["q3"]) == (6000, 6000, None)
    toner = by_name["หมึกพิมพ์"]  # the plan's own name and unit (D-43)
    assert (toner["item_id"], toner["unit"], toner["note"]) == (None, "box", None)
    glove = by_name["ถุงมือ"]
    assert (glove["plan_type"], glove["purchasing_department_id"]) == ("CENTRAL", "MOCK-DEP-STORE")
    ((_, _, after),) = plan.audit("PLAN_IMPORTED")
    assert after["rows_imported"] == 4 and "pre_bound" not in after
    assert after["file_sha256"] == body["file_sha256"]

    again = do_import(plan, GOOD)
    assert again.status_code == 409 and again.json()["detail"]["code"] == "DUPLICATE_FILE"
    dup = preview(plan, GOOD).json()
    assert not dup["can_import"] and dup["duplicate_of"]["id"] == imp["id"]
    changed = do_import(plan, workbook([H_ENT], GOOD_ROWS[:1]), expected_sha256=body["file_sha256"])
    assert changed.status_code == 409 and changed.json()["detail"]["code"] == "FILE_CHANGED"
    listed = hx.client.get(f"/api/plan-years/{FY}/imports", headers=plan.h).json()
    assert [x["id"] for x in listed] == [imp["id"]]


@pytest.mark.spec("D-42", "AT-40.17.2")
def test_any_error_imports_nothing_and_every_error_is_reported(plan: Plan) -> None:
    headings = [
        H_ENT,
        H_ENT,  # duplicate heading code
        ["H3", "", "ASSIGNED", "MOCK-DEP-ENT", "MOCK-DEP-STORE", "MOCK-FUND-1", "MOCK-CAT-OFFICE"],
        ["H4", "", "CENTRAL", "MOCK-DEP-ENT", None, "MOCK-FUND-1", "MOCK-CAT-OFFICE"],
        ["H5", "", "DEPARTMENT", "MOCK-DEP-OLD", None, "MOCK-FUND-2", "MOCK-CAT-OFFICE"],
        ["H6", "", "DEPARTMENT", "NO-SUCH-DEP", None, "MOCK-FUND-1", "MOCK-CAT-OFFICE"],
        ["H7", "", "SOMETHING", "MOCK-DEP-ENT", None, "MOCK-FUND-1", "MOCK-CAT-OFFICE"],
        ["H8", "", "DEPARTMENT", "MOCK-DEP-IT", None, "MOCK-FUND-1", "MOCK-CAT-OFFICE"],
        ["H9", "", "DEPARTMENT", "MOCK-DEP-IT", None, "MOCK-FUND-1", "MOCK-CAT-MED"],
    ]
    rows = [
        row("H8", "ok row", "box", 1, 1, 1),
        row("H8", "bad qty", "box", "สิบ", 1, 10),
        row("H8", "negative", "box", -1, 1, -1),
        row("H8", "too precise", "box", 1, "0.001", "0.001"),
        row("H8", "zero qty", "box", 0, 1, 5),
        row("H8", "", "box", 1, 1, 1),
        row("HX", "unknown heading", "box", 1, 1, 1),
    ]
    data = workbook(headings, rows)
    body = preview(plan, data).json()
    assert not body["can_import"] and body["rows_to_import"] == 0
    found = {(e["code"], e["sheet"], e["row"]) for e in body["errors"]}
    expected = {
        ("DUPLICATE_HEADING", "หัวข้อ", 2),
        ("DUPLICATE_HEADING", "หัวข้อ", 3),
        ("PLAN_TYPE_NOT_IMPORTED", "หัวข้อ", 4),
        ("CENTRAL_WITHOUT_PURCHASER", "หัวข้อ", 5),
        ("INACTIVE_CODE", "หัวข้อ", 6),
        ("NO_BUDGET_LINE", "หัวข้อ", 6),
        ("UNKNOWN_CODE", "หัวข้อ", 7),
        ("UNKNOWN_PLAN_TYPE", "หัวข้อ", 8),
        ("BAD_NUMBER", "รายการ", 3),
        ("NEGATIVE", "รายการ", 4),
        ("TOO_PRECISE", "รายการ", 5),
        ("ZERO_QTY_WITH_AMOUNT", "รายการ", 6),
        ("MISSING_VALUE", "รายการ", 7),
        ("UNKNOWN_HEADING", "รายการ", 8),
    }
    assert expected <= found, sorted(expected - found)
    r = do_import(plan, data)
    assert r.status_code == 409 and r.json()["detail"]["code"] == "IMPORT_HAS_ERRORS"
    assert r.json()["detail"]["details"]
    assert count(plan.hx.engine, "SELECT count(*) FROM plan_item") == 0
    assert count(plan.hx.engine, "SELECT count(*) FROM plan_import") == 0
    assert plan.audit("PLAN_IMPORTED") == []


@pytest.mark.spec("D-42", "AT-40.17.2")
def test_a_file_that_is_not_the_template_is_refused(plan: Plan) -> None:
    for data, code in (
        (b"not an excel file", "NOT_XLSX"),
        (workbook([H_ENT], [], drop="รายการ"), "MISSING_SHEET"),
    ):
        r = preview(plan, data)
        assert r.status_code == 422 and r.json()["detail"]["code"] == code, r.text
    wb = Workbook()
    ws = wb.active
    ws.title = "หัวข้อ"
    ws.append(["รหัสหัวข้อ"])
    wb.create_sheet("รายการ").append(ROW_TITLES)
    out = BytesIO()
    wb.save(out)
    r = preview(plan, out.getvalue())
    assert r.status_code == 422 and r.json()["detail"]["code"] == "MISSING_COLUMNS"
    big = post(plan, f"/api/plan-years/{FY}/imports/preview", b"x" * (10 * 1024 * 1024 + 1))
    assert big.status_code == 422 and big.json()["detail"]["code"] == "FILE_TOO_LARGE"
    empty = post(plan, f"/api/plan-years/{FY}/imports/preview", b"")
    assert empty.status_code == 422 and empty.json()["detail"]["code"] == "EMPTY_FILE"


# ------------------------------------------------------------------ rows without a code
@pytest.mark.spec("D-42", "D-43", "AT-40.17.4", "AT-40.18.1")
def test_imported_rows_are_activated_and_chosen_by_the_requester(plan: Plan) -> None:
    hx = plan.hx
    assert do_import(plan, GOOD).status_code == 201
    # Unbound rows are no duplicates of each other or of a manual item of another code.
    v = hx.client.get(f"/api/plan-years/{FY}/validation", headers=plan.h).json()
    assert v["can_activate"], v
    assert plan.move("APPROVED").status_code == 200
    assert plan.move("ACTIVE").status_code == 200
    with hx.engine.connect() as conn:
        n, qty = conn.execute(
            text(
                "SELECT count(*), sum(qty_delta) FROM plan_ledger "
                "WHERE event_type = 'PLAN_ACTIVATED'"
            )
        ).one()
    assert (n, qty) == (4, Decimal("157"))
    # The Central Plan usage report: an unbound central row has no issue figures - shown as
    # unknown with a note, never as 0.
    from datetime import date
    from types import MappingProxyType

    from ppr.integration.hosxp.contracts import StockMovementSource, StockVerification

    hx.gateway.data.stock_source = StockMovementSource(
        problem=None,
        verification=StockVerification("MOCK IT", date(2026, 10, 1), "MOCK-REF"),
        kind_map=MappingProxyType({}),
    )
    rep = hx.client.get("/api/reports/CENTRAL_PLAN_USAGE", headers=plan.h)
    assert rep.status_code == 200, rep.text
    rows = rep.json()["rows"]
    gloves = [r for r in rows if r["item_name"] in ("ถุงมือ", "ถุงมือ ไซซ์พิเศษ")]
    assert len(gloves) == 2
    for g in gloves:
        assert g["item_code"] is None and g["note"].startswith("ไม่มีรหัสสินค้า HOSxP")
        assert g["issued_qty"] == "ยังไม่มีข้อมูล" and g["net_used_qty"] == "ยังไม่มีข้อมูล"
    req = hx.user("req", Role.REQUESTER)
    hx.put_pr("PR-IMP-1", ("MOCK-ITEM-TONER", "2", "1500"))
    # never matched by itself: the requester sees the rows counted in boxes and chooses one
    first = hx.client.get("/api/prs/PR-IMP-1/prevalidation", headers=hx.h(req)).json()
    assert not first["eligible"]
    assert [v["code"] for v in first["lines"][0]["violations"]] == ["PLAN_ROW_NOT_CHOSEN"]
    [option] = first["options"]
    assert (option["item_name"], option["unit"], option["item_id"]) == ("หมึกพิมพ์", "box", None)
    ok = hx.client.get(
        "/api/prs/PR-IMP-1/prevalidation",
        params={"choice": f"PR-IMP-1-1:{option['id']}"},
        headers=hx.h(req),
    ).json()
    assert ok["eligible"], ok
    assert ok["lines"][0]["matched_plan_item_id"] == option["id"]
    hx.put_pr("PR-IMP-2", ("MOCK-ITEM-PAPER", "2", "120"))
    no = hx.client.get("/api/prs/PR-IMP-2/prevalidation", headers=hx.h(req)).json()
    assert not no["eligible"]
    assert [v["code"] for v in no["lines"][0]["violations"]] == ["NOT_IN_PLAN"]
    # the year is ACTIVE: no import, no removal
    r = do_import(plan, workbook([H_ENT], [row("H1", "late", "box", 1, 1, 1)]))
    assert r.status_code == 409 and r.json()["detail"]["code"] == "PLAN_YEAR_NOT_DRAFT"
    imp = hx.client.get(f"/api/plan-years/{FY}/imports", headers=plan.h).json()[0]
    rm = hx.client.post(
        f"/api/plan-imports/{imp['id']}/remove", json={"reason": "x"}, headers=plan.h
    )
    assert rm.status_code == 409 and rm.json()["detail"]["code"] == "PLAN_YEAR_NOT_DRAFT"
    # a row without a HOSxP item is not given one by amendment (D-43)
    items = {
        i["item_name"]: i
        for i in hx.client.get(f"/api/plan-years/{FY}/items", headers=plan.h).json()
    }
    paper = items["กระดาษ A4 80 แกรม"]
    am = hx.client.post(
        f"/api/plan-years/{FY}/amendments/preview",
        json={
            "approval_document_no": "D1",
            "approval_date": "2026-10-07",
            "reason": "bind",
            "items": [{"plan_item_id": paper["id"], "item_id": "MOCK-ITEM-PAPER"}],
        },
        headers=plan.h,
    )
    assert am.status_code == 422 and am.json()["detail"]["code"] == "UNBOUND_ROW", am.text


@pytest.mark.spec("D-42", "D-43", "AT-40.17.4", "AT-40.18.5")
def test_an_imported_row_is_edited_in_draft_and_given_an_item_only_in_its_unit(
    plan: Plan,
) -> None:
    hx = plan.hx
    assert do_import(plan, GOOD).status_code == 201
    items = {
        i["item_name"]: i
        for i in hx.client.get(f"/api/plan-years/{FY}/items", headers=plan.h).json()
    }
    paper = items["กระดาษ A4 80 แกรม"]
    body = {
        "plan_budget_id": paper["plan_budget_id"],
        "plan_type": "DEPARTMENT",
        "owner_department_id": "MOCK-DEP-ENT",
        "planned_qty": "90",
        "estimated_unit_price": "120",
        "planned_amount": "10800",
    }
    r = hx.client.put(f"/api/plan-items/{paper['id']}", json=body, headers=plan.h)
    assert r.status_code == 200, r.text
    assert (r.json()["item_id"], r.json()["item_name"], r.json()["unit"]) == (
        None,
        "กระดาษ A4 80 แกรม",
        "รีม",
    )
    wrong_unit = hx.client.put(
        f"/api/plan-items/{paper['id']}",
        json={**body, "item_id": "MOCK-ITEM-PAPER"},
        headers=plan.h,
    )
    assert wrong_unit.status_code == 422 and wrong_unit.json()["detail"]["code"] == "UNIT_MISMATCH"
    # D-43: a manual item may have no HOSxP item; it then needs its own name and unit
    manual = plan.item(paper["plan_budget_id"], "100", item_id=None)
    assert manual.status_code == 422
    assert manual.json()["detail"]["code"] == "NAME_AND_UNIT_REQUIRED"
    named = plan.item(
        paper["plan_budget_id"], "100", item_id=None, item_name=" แฟ้มเอกสาร ", unit="แฟ้ม"
    )
    assert named.status_code == 201, named.text
    assert (named.json()["item_id"], named.json()["item_name"], named.json()["unit"]) == (
        None,
        "แฟ้มเอกสาร",
        "แฟ้ม",
    )
    assigned = plan.item(
        paper["plan_budget_id"],
        "100",
        item_id=None,
        item_name="x",
        unit="y",
        plan_type="ASSIGNED",
        purchasing_department_id="MOCK-DEP-STORE",
    )
    assert assigned.status_code == 422 and assigned.json()["detail"]["code"] == "ITEM_REQUIRED"


# ------------------------------------------------------------------ removal
@pytest.mark.spec("D-42", "AT-40.17.5")
def test_an_import_is_removed_with_a_reason_while_draft(plan: Plan) -> None:
    hx = plan.hx
    imp = do_import(plan, GOOD).json()
    path = f"/api/plan-imports/{imp['id']}/remove"
    assert hx.client.post(path, json={"reason": " "}, headers=plan.h).status_code == 422
    r = hx.client.post(path, json={"reason": "ผิดไฟล์"}, headers=plan.h)
    assert r.status_code == 200, r.text
    assert (r.json()["state"], r.json()["removed_reason"]) == ("REMOVED", "ผิดไฟล์")
    assert count(hx.engine, "SELECT count(*) FROM plan_item") == 0
    assert count(hx.engine, "SELECT count(*) FROM plan_import") == 1  # the record stays
    ((_, reason, after),) = plan.audit("PLAN_IMPORT_REMOVED")
    assert reason == "ผิดไฟล์" and after["plan_items_deleted"] == 4
    again = hx.client.post(path, json={"reason": "again"}, headers=plan.h)
    assert again.status_code == 409 and again.json()["detail"]["code"] == "IMPORT_ALREADY_REMOVED"
    # the same file may be imported again once its import is removed
    assert do_import(plan, GOOD).status_code == 201


@pytest.mark.spec("D-42")
def test_only_planning_imports(hx: Harness, plan: Plan) -> None:
    req = hx.user("req2", Role.REQUESTER)
    r = TestClient(hx.client.app).post(
        f"/api/plan-years/{FY}/imports",
        content=GOOD,
        headers={**hx.h(req), "Content-Type": "application/octet-stream"},
    )
    assert r.status_code == 403


@pytest.mark.spec("D-42", "AT-40.17.4")
def test_every_plan_view_and_report_shows_unbound_rows(plan: Plan) -> None:
    hx = plan.hx
    assert do_import(plan, GOOD).status_code == 201
    assert plan.move("APPROVED").status_code == 200
    assert plan.move("ACTIVE").status_code == 200
    balances = hx.client.get(f"/api/plan-years/{FY}/balances", headers=plan.h)
    assert balances.status_code == 200, balances.text
    unbound = [b for b in balances.json() if b["item_id"] is None]
    assert len(unbound) == 4 and all(b["item_code"] is None for b in unbound)
    assert hx.client.get("/api/dashboard", headers=plan.h).status_code == 200
    assert hx.client.get("/api/alerts", headers=plan.h).status_code == 200
    codes = [r["code"] for r in hx.client.get("/api/reports", headers=plan.h).json()]
    for code in codes:
        r = hx.client.get(f"/api/reports/{code}", headers=plan.h, params={"fiscal_year": FY})
        assert r.status_code == 200, (code, r.text)
        x = hx.client.get(f"/api/reports/{code}/xlsx", headers=plan.h, params={"fiscal_year": FY})
        assert x.status_code == 200, (code, x.text)


@pytest.mark.spec("D-42", "AT-40.17.2")
def test_numbers_beyond_the_columns_and_loose_commas_are_errors(plan: Plan) -> None:
    rows = [
        row("H1", "huge amount", "box", 1, 1, 1e20),
        row("H1", "huge qty", "box", 1e16, 0, 0),
        row("H1", "comma", "box", "1,5", 1, 1),
        row("H1", "formula result", "box", 3, 1.1, 3 * 1.1),  # 3.3000000000000003 in Excel
    ]
    body = preview(plan, workbook([H_ENT], rows)).json()
    found = {(e["code"], e["row"]) for e in body["errors"]}
    assert found == {("TOO_PRECISE", 2), ("TOO_PRECISE", 3), ("BAD_NUMBER", 4)}, body["errors"]


@pytest.mark.spec("D-43", "AT-40.18.5")
def test_an_item_column_from_the_old_template_is_not_read(plan: Plan) -> None:
    wb = Workbook()
    wb.active.title = "คำอธิบาย"
    wb.create_sheet("หัวข้อ").append(HEAD_TITLES)
    wb["หัวข้อ"].append(H_ENT)
    rows = wb.create_sheet("รายการ")
    rows.append([*ROW_TITLES, "รหัสสินค้า HOSxP"])
    rows.append([*row("H1", "หมึก", "box", 1, 18000, 18000), "MOCK-ITEM-TONER"])
    out = BytesIO()
    wb.save(out)
    body = preview(plan, out.getvalue()).json()
    assert body["can_import"], body
    assert [w["code"] for w in body["warnings"]] == ["ITEM_COLUMN_IGNORED"]
    assert "รหัสสินค้า HOSxP" not in ROW_TITLES  # the template has no item column
    assert do_import(plan, out.getvalue()).status_code == 201
    [item] = plan.hx.client.get(f"/api/plan-years/{FY}/items", headers=plan.h).json()
    assert (item["item_id"], item["item_name"], item["unit"]) == (None, "หมึก", "box")


@pytest.mark.spec("D-42", "AT-40.17.4", "AT-40.17.5")
def test_an_unbound_row_is_amended_in_value_and_its_history_kept(plan: Plan) -> None:
    hx = plan.hx
    assert do_import(plan, GOOD).status_code == 201
    assert plan.move("APPROVED").status_code == 200
    assert plan.move("ACTIVE").status_code == 200
    items = {
        i["item_name"]: i
        for i in hx.client.get(f"/api/plan-years/{FY}/items", headers=plan.h).json()
    }
    paper, glove = items["กระดาษ A4 80 แกรม"], items["ถุงมือ"]
    r = hx.client.post(
        f"/api/plan-years/{FY}/amendments",
        json={
            "approval_document_no": "สส 1/2570",
            "approval_date": "2026-10-07",
            "reason": "ย้ายยอด",
            "items": [
                {"plan_item_id": paper["id"], "planned_qty": "90", "planned_amount": "10800"},
                {
                    "plan_item_id": items["หมึกพิมพ์"]["id"],
                    "planned_qty": "4.8",
                    "planned_amount": "7200",
                },
            ],
        },
        headers=plan.h,
    )
    assert r.status_code == 201, r.text
    assert glove["item_id"] is None
    rep = hx.client.get("/api/reports/PLAN_AMENDMENT", headers=plan.h, params={"fiscal_year": FY})
    assert rep.status_code == 200, rep.text
    shown = [r for r in rep.json()["rows"] if r.get("item_name") == "กระดาษ A4 80 แกรม"]
    assert shown and all(r["item_code"] in ("", None) for r in shown)
    assert "None" not in str(rep.json()["rows"])


@pytest.mark.spec("D-42", "AT-40.17.5")
def test_the_import_audit_stays_after_removal(plan: Plan) -> None:
    imp = do_import(plan, GOOD).json()
    r = plan.hx.client.post(
        f"/api/plan-imports/{imp['id']}/remove", json={"reason": "ผิด"}, headers=plan.h
    )
    assert r.status_code == 200
    assert len(plan.audit("PLAN_IMPORTED")) == 1 and len(plan.audit("PLAN_IMPORT_REMOVED")) == 1
