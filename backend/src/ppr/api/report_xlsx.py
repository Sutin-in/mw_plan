"""Excel (.xlsx) file of a report (D-35). Same rows as the screen; numbers stay numbers."""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ppr.application.report_services import HOSPITAL_TZ, ColType, Report

_FORMATS = {
    ColType.MONEY: "#,##0.00",
    ColType.QTY: "#,##0.00##",
    ColType.PERCENT: "0.00",
    ColType.INT: "0",
    ColType.DATETIME: "yyyy-mm-dd hh:mm",
    ColType.DATE: "yyyy-mm-dd",
}
# A text cell starting with one of these could be read as a formula by a spreadsheet.
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _value(v: Any) -> Any:
    if isinstance(v, Decimal):
        return float(v) if v != v.to_integral_value() or abs(v) > 2**53 else int(v)
    if isinstance(v, datetime):
        return v.astimezone(HOSPITAL_TZ).replace(tzinfo=None)  # hospital time, as on screen
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)  # control characters cannot be stored
    return v


def _put(ws: Any, row: int, col: int, value: Any, fmt: str | None = None) -> None:
    cell = ws.cell(row=row, column=col)
    cell.value = _value(value)
    if isinstance(cell.value, str) and cell.value.startswith(_FORMULA_START):
        cell.data_type = "s"  # always text, never a formula
    if fmt and cell.value is not None:
        cell.number_format = fmt


def build(report: Report) -> bytes:
    d = report.definition
    wb = Workbook()
    ws = wb.active
    ws.title = d.code[:31]
    _put(ws, 1, 1, d.title)
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    _put(ws, 2, 1, "จัดทำเมื่อ")
    _put(ws, 2, 2, report.generated_at, _FORMATS[ColType.DATETIME])
    r = 3
    for label, value in report.filters:
        _put(ws, r, 1, label)
        _put(ws, r, 2, value)
        r += 1
    for note in report.notes:
        _put(ws, r, 1, note)
        r += 1
    r += 1
    header = r
    for c, col in enumerate(d.columns, 1):
        _put(ws, header, c, col.label)
        ws.cell(row=header, column=c).font = Font(bold=True)
    for i, row in enumerate(report.rows, 1):
        for c, col in enumerate(d.columns, 1):
            _put(ws, header + i, c, row.get(col.key), _FORMATS.get(col.type))
    for c, col in enumerate(d.columns, 1):
        width = 14 if col.type is not ColType.TEXT else max(12, min(40, len(col.label) + 6))
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = ws.cell(row=header + 1, column=1)
    if report.rows:
        last = get_column_letter(len(d.columns))
        ws.auto_filter.ref = f"A{header}:{last}{header + len(report.rows)}"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
